"""Exporta um relatório do livro (`report_book`) em Excel (.xlsx) ou CSV.

Excel: aba "Resumo" (título, período, filial, filtros aplicados e indicadores com a variação)
e uma aba por bloco (tabela, ranking, participação, série no tempo, mapa de calor, cascata),
com moeda/percentual/data formatados, cabeçalho fixo e filtro do Excel. Tabela com abas (Curva
ABC por classe) ganha também uma aba por classe. `bloco` exporta só um bloco; `aba` só uma classe.

CSV (abre no Excel brasileiro): ponto e vírgula, vírgula decimal e BOM UTF-8; um bloco só
(o pedido, ou a primeira tabela do relatório).
"""
import csv
import io
import re
from datetime import date
from decimal import Decimal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

FORMATS = {"money": '"R$" #,##0.00;[Red]-"R$" #,##0.00', "percent": "0.0%", "quantity": "#,##0.###",
           "number": "#,##0", "decimal": "#,##0.0", "date": "DD/MM/YYYY"}
HEAD_FILL = PatternFill("solid", fgColor="E8ECF2")
TITLE_FONT = Font(bold=True, size=14)
BOLD = Font(bold=True)


def _parse_date(value):
    if isinstance(value, date):
        return value
    if isinstance(value, str) and re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
        return date.fromisoformat(value)
    return None


def _cell_value(value, fmt):
    """Valor para a planilha: número de verdade (não texto), % como fração, data como data."""
    if value is None or value == "":
        return None
    if fmt == "percent":
        return float(Decimal(str(value)) / 100)
    if fmt in ("money", "quantity", "number", "decimal"):
        try:
            return float(Decimal(str(value)))
        except ArithmeticError:
            return str(value)
    if fmt == "date":
        return _parse_date(value) or str(value)
    return value if isinstance(value, (int, float)) else str(value)


# ── blocos → (colunas, linhas) ──────────────────────────────────────────────
def block_table(block, tab=None):
    """Qualquer bloco vira uma tabela: [(rótulo, chave, formato)], [linhas]."""
    kind, fmt = block["type"], block.get("format", "number")
    if kind == "table":
        rows = block["rows"]
        tabs = block.get("tabs")
        if tabs and tab:
            rows = [r for r in rows if str(r.get(tabs["key"])) == str(tab)]
        return [(c["label"], c["key"], c.get("format", "text")) for c in block["columns"]], rows
    if kind == "ranking":
        cols = [("#", "pos", "number"), ("Item", "label", "text"), ("Valor", "value", fmt), ("Detalhe", "sub", "text")]
        if any(r.get("group") for r in block["rows"]):
            cols.insert(2, ("Grupo", "group", "text"))
        return cols, [{**r, "pos": i} for i, r in enumerate(block["rows"], start=1)]
    if kind == "share":
        total = sum((Decimal(str(r["value"] or 0)) for r in block["rows"]), Decimal("0"))
        return [("Item", "label", "text"), ("Valor", "value", fmt), ("Participação", "pct", "percent")], [
            {**r, "pct": (Decimal(str(r["value"] or 0)) / total * 100) if total else 0} for r in block["rows"]]
    if kind in ("line", "bars"):
        series = block["series"]
        grain = block.get("grain")
        if grain:  # semana/mês por extenso ("19/10 a 25/10/2026", "Outubro/2026"); dia como data de verdade
            from starrp.services.report_book import period_label

            xs = [_parse_date(x) if grain == "day" else period_label(_parse_date(x), grain) for x in block["x"]]
            cols = [("Período", "x", "date" if grain == "day" else "text")]
            cols += [(s["name"], f"s{i}", fmt) for i, s in enumerate(series)]
            return cols, [{"x": x, **{f"s{i}": s["values"][j] for i, s in enumerate(series)}} for j, x in enumerate(xs)]
        cols = [("Período", "x", "date" if _parse_date(block["x"][0] if block["x"] else None) else "text")]
        cols += [(s["name"], f"s{i}", fmt) for i, s in enumerate(series)]
        return cols, [{"x": x, **{f"s{i}": s["values"][j] for i, s in enumerate(series)}} for j, x in enumerate(block["x"])]
    if kind == "heatmap":
        cols = [("", "row", "text")] + [(c, f"c{i}", fmt) for i, c in enumerate(block["cols"])]
        return cols, [{"row": r, **{f"c{i}": v for i, v in enumerate(block["values"][k])}} for k, r in enumerate(block["rows"])]
    if kind == "compare":
        return [("", "label", "text")] + [(m["label"], m["key"], m.get("format", "number")) for m in block["metrics"]], block["rows"]
    if kind == "insights":
        tones = {"good": "Positivo", "bad": "Atenção", "info": "Observação"}
        return [("Conclusão", "text", "text"), ("Tipo", "tone_label", "text")], [
            {**i, "tone_label": tones.get(i.get("tone"), "")} for i in block["items"]]
    if kind == "waterfall":
        kinds = {"add": "Soma", "sub": "Desconta", "total": "Subtotal"}
        return [("Linha", "label", "text"), ("Valor", "value", fmt), ("Tipo", "kind_label", "text")], [
            {**s, "kind_label": kinds.get(s["kind"], s["kind"])} for s in block["steps"]]
    return [], []


def _sheet_name(title, used):
    """Nome de aba do Excel (até 31 letras): corta numa palavra inteira e põe "…", nunca no meio."""
    name = re.sub(r"\s+", " ", re.sub(r"[\[\]:*?/\\]", " ", title)).strip() or "Dados"
    if len(name) > 31:
        name = f"{name[:30].rsplit(' ', 1)[0].rstrip(' -(.')}…"
    base, n = name, 2
    while name.lower() in used:
        suffix = f" ({n})"
        name, n = f"{base[:31 - len(suffix)]}{suffix}", n + 1
    used.add(name.lower())
    return name


def _write_table(ws, cols, rows, start_row=1):
    for c, (label, _, _) in enumerate(cols, start=1):
        cell = ws.cell(row=start_row, column=c, value=label)
        cell.font, cell.fill = BOLD, HEAD_FILL
        cell.alignment = Alignment(horizontal="right" if cols[c - 1][2] in FORMATS and cols[c - 1][2] != "date" else "left")
    for r, row in enumerate(rows, start=start_row + 1):
        for c, (_, key, fmt) in enumerate(cols, start=1):
            cell = ws.cell(row=r, column=c, value=_cell_value(row.get(key), fmt))
            if fmt in FORMATS:
                cell.number_format = FORMATS[fmt]
            if row.get("strong"):
                cell.font = BOLD
    for c, (label, key, fmt) in enumerate(cols, start=1):
        sample = [len(str(label))] + [len(str(row.get(key) or "")) + (4 if fmt == "money" else 0) for row in rows[:300]]
        ws.column_dimensions[get_column_letter(c)].width = min(60, max(8, max(sample) + 2))
    if rows:
        ws.auto_filter.ref = f"A{start_row}:{get_column_letter(len(cols))}{start_row + len(rows)}"
    ws.freeze_panes = ws.cell(row=start_row + 1, column=1)


def _summary(ws, report, context):
    ws["A1"] = report["title"]
    ws["A1"].font = TITLE_FONT
    ws["A2"] = report["subtitle"]
    period = report["period"]
    lines = [("Período", f"{_parse_date(period['from']):%d/%m/%Y} a {_parse_date(period['to']):%d/%m/%Y}"),
             ("Filial", context.get("store") or "Todas as filiais permitidas"),
             ("Filtros", context.get("filters") or "nenhum"),
             ("Emitido por", context.get("user", "")), ("Emitido em", context.get("issued", ""))]
    for i, (label, value) in enumerate(lines, start=4):
        ws.cell(row=i, column=1, value=label).font = BOLD
        ws.cell(row=i, column=2, value=value)
    start = 4 + len(lines) + 1
    cols = [("Indicador", "label", "text"), ("Valor", "value", None), ("Variação vs. período anterior", "delta", "percent"),
            ("Observação", "hint", "text")]
    for c, (label, _, _) in enumerate(cols, start=1):
        cell = ws.cell(row=start, column=c, value=label)
        cell.font, cell.fill = BOLD, HEAD_FILL
    for r, k in enumerate(report["kpis"], start=start + 1):
        ws.cell(row=r, column=1, value=k["label"])
        value = ws.cell(row=r, column=2, value=_cell_value(k["value"], k["format"]))
        if k["format"] in FORMATS:
            value.number_format = FORMATS[k["format"]]
        delta = ws.cell(row=r, column=3, value=_cell_value(k["delta"], "percent"))
        delta.number_format = '+0.0%;-0.0%;0.0%'
        ws.cell(row=r, column=4, value=k.get("hint") or "")
    for col, width in zip("ABCD", (34, 22, 30, 60), strict=True):
        ws.column_dimensions[col].width = width


def to_xlsx(report, context, block_index=None, tab=None):
    wb = Workbook()
    used = set()
    blocks = list(enumerate(report["blocks"]))
    if block_index is not None:
        blocks = [blocks[block_index]]
        wb.remove(wb.active)
    else:
        ws = wb.active
        ws.title = _sheet_name("Resumo", used)
        _summary(ws, report, context)
    for _, block in blocks:
        cols, rows = block_table(block, tab)
        if not cols:
            continue
        title = block["title"] + (f" - {block['tabs']['label']} {tab}" if tab and block.get("tabs") else "")
        _write_table(wb.create_sheet(_sheet_name(title, used)), cols, rows)
        tabs = block.get("tabs")
        if tabs and not tab and block_index is None:  # uma aba por classe (A, B, C...)
            for value in sorted({str(r.get(tabs["key"])) for r in block["rows"]}):
                sub_cols, sub_rows = block_table(block, value)
                _write_table(wb.create_sheet(_sheet_name(f"{tabs['label']} {value}", used)), sub_cols, sub_rows)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _csv_number(value, fmt):
    """Número com vírgula e as casas certas: R$ com 2, % com 1, quantidade sem zeros sobrando."""
    number = Decimal(str(value))
    if fmt == "money":
        number = number.quantize(Decimal("0.01"))
    elif fmt in ("percent", "decimal"):
        number = number.quantize(Decimal("0.1"))
    elif fmt == "quantity":
        number = number.quantize(Decimal("0.001")).normalize()
    return f"{number:f}".replace(".", ",")


def to_csv(report, block_index=None, tab=None):
    blocks = report["blocks"]
    if block_index is None:
        block_index = next((i for i, b in enumerate(blocks) if b["type"] == "table"), 0)
    cols, rows = block_table(blocks[block_index], tab)
    out = io.StringIO()
    writer = csv.writer(out, delimiter=";", quoting=csv.QUOTE_ALL)
    writer.writerow([label for label, _, _ in cols])
    for row in rows:
        line = []
        for _, key, fmt in cols:
            value = row.get(key)
            if value is None:
                line.append("")
            elif fmt == "date" and _parse_date(value):
                line.append(f"{_parse_date(value):%d/%m/%Y}")
            elif fmt in ("money", "percent", "quantity", "number", "decimal"):
                line.append(_csv_number(value, fmt))
            else:
                line.append(str(value))
        writer.writerow(line)
    return ("﻿" + out.getvalue()).encode("utf-8")


def filters_label(report, params):
    """"Vendedor: Carla Mendes · Canal: PDV" — os filtros aplicados, com o nome (não o id)."""
    parts = []
    for flt in report.get("filters", []):
        value = params.get(flt["name"])
        if value:
            label = next((o["label"] for o in flt["options"] if o["value"] == str(value)), value)
            parts.append(f"{flt['label']}: {label}")
    return " · ".join(parts)


def filename(slug, report, ext, tab=None):
    period = report["period"]
    suffix = f"-{tab}" if tab else ""
    return f"{slug}{suffix}-{period['from']}-a-{period['to']}.{ext}"
