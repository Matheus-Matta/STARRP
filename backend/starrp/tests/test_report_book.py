"""Livro de relatórios: todos respondem no formato padrão e os números batem com a venda."""
from decimal import Decimal

import pytest

from starrp.services.report_book import BOOK

from .test_sale_flow import _stock


@pytest.fixture
def sold(ctx):
    ctx.account.extra_modules = ["stock", "finance"]
    ctx.account.save()
    product = ctx.product(name="Mesa", price="50.00")
    _stock(ctx, product, 10, cost="20.00")
    api = ctx.client()
    sale = api.post("/api/v1/sales/", {}, format="json").json()
    api.post(f"/api/v1/sales/{sale['id']}/items/", {"variant": str(product.pk), "quantity": "2"}, format="json")
    api.post(f"/api/v1/sales/{sale['id']}/checkout/",
             {"payments": [{"nsu": "000123", "method": str(ctx.methods["pix"].pk), "amount": "100.00"}]}, format="json")
    return ctx, api


@pytest.mark.parametrize("slug", list(BOOK))
def test_todo_relatorio_responde_no_formato_padrao(sold, slug):
    _, api = sold
    body = api.get(f"/api/v1/reports/book/{slug}/").json()

    assert body.get("title") and body["kpis"] and body["blocks"], body
    for block in body["blocks"]:
        assert block["type"] in {"line", "bars", "ranking", "share", "heatmap", "waterfall", "table", "compare", "insights"}


def test_vendas_e_curva_abc_batem_com_a_venda(sold):
    _, api = sold
    sales = api.get("/api/v1/reports/book/vendas/").json()
    abc = api.get("/api/v1/reports/book/curva-abc/").json()

    assert Decimal(sales["kpis"][0]["value"]) == Decimal("100.00") and sales["kpis"][0]["delta"] is None
    assert Decimal(sales["kpis"][3]["value"]) == Decimal("60.00")
    table = next(b for b in abc["blocks"] if b["type"] == "table")
    assert table["rows"][0]["abc"] == "A" and Decimal(table["rows"][0]["cumulative"]) == 100


def test_relatorio_contabil_exige_permissao_e_modulo(sold):
    ctx, _ = sold
    seller = ctx.user("vend_rel", "seller")
    denied = ctx.client(seller).get("/api/v1/reports/book/dre/")
    unknown = ctx.client().get("/api/v1/reports/book/nao-existe/")

    assert denied.status_code == 403 and unknown.status_code == 404


def test_exporta_excel_com_resumo_e_csv_de_um_bloco(sold):
    import io

    import openpyxl

    _, api = sold
    xlsx = api.get("/api/v1/reports/book/curva-abc/export/", {"formato": "xlsx"})
    book = openpyxl.load_workbook(io.BytesIO(xlsx.content))
    blocks = api.get("/api/v1/reports/book/curva-abc/").json()["blocks"]
    table = next(i for i, b in enumerate(blocks) if b["type"] == "table")
    csv_b = api.get("/api/v1/reports/book/curva-abc/export/", {"formato": "csv", "bloco": str(table), "aba": "A"})
    lines = csv_b.content.decode("utf-8-sig").splitlines()

    assert xlsx.status_code == 200 and xlsx["Content-Disposition"].endswith('.xlsx"')
    assert book.sheetnames[0] == "Resumo" and "Classe A" in book.sheetnames
    assert book["Resumo"]["A1"].value == "Curva ABC de produtos"
    assert csv_b.status_code == 200 and "classe-a" in csv_b["Content-Disposition"]
    assert lines[0].startswith('"#";"Produto"') and '"100,00"' in lines[1]  # R$ 100,00 com 2 casas


def test_cruzamento_cliente_abc_com_avaliacao(sold):
    """A nota que o cliente deu aparece na classe dele (cliente A dá nota 5) e vira conclusão escrita."""
    from starrp.models import Customer, Review, Sale

    ctx, api = sold
    customer = Customer.all_objects.create(account=ctx.account, name="Bruna Lima")
    sale = Sale.all_objects.filter(account=ctx.account, status="completed").first()
    Sale.all_objects.filter(pk=sale.pk).update(customer=customer)
    product = sale.items.first().variant.product
    Review.all_objects.create(account=ctx.account, store=ctx.store, kind="product", product=product, sale=sale,
                              customer=customer, public_name="Bruna L.", rating=5, status="approved")

    body = api.get("/api/v1/reports/book/cruzamento-clientes/").json()
    grid = next(b for b in body["blocks"] if b["type"] == "compare")
    notes = next(b for b in body["blocks"] if b["type"] == "insights")

    row = next(r for r in grid["rows"] if r["label"] == "Classe A")
    assert row["customers"] == 1 and Decimal(row["rating"]) == 5
    assert any("Base ainda pequena" in i["text"] for i in notes["items"])


def test_cupom_usado_aparece_com_desconto_lucro_e_veredito(ctx):
    """Venda com cupom de 10%: o relatório de cupons mostra o uso, o desconto e um veredito."""
    from starrp.models import Coupon, Customer

    ctx.account.extra_modules = ["stock"]
    ctx.account.save()
    Coupon.all_objects.create(account=ctx.account, code="VOLTE10", kind="percent", value="10.00")
    customer = Customer.all_objects.create(account=ctx.account, name="Rita Souza")
    product = ctx.product(name="Cadeira", price="100.00")
    _stock(ctx, product, 5, cost="40.00")
    api = ctx.client()
    sale = api.post("/api/v1/sales/", {"customer": str(customer.pk)}, format="json").json()
    api.post(f"/api/v1/sales/{sale['id']}/items/", {"variant": str(product.pk)}, format="json")
    applied = api.post(f"/api/v1/sales/{sale['id']}/coupon/", {"code": "VOLTE10"}, format="json").json()
    api.post(f"/api/v1/sales/{sale['id']}/checkout/", {"payments": [
        {"nsu": "000777", "method": str(ctx.methods["pix"].pk), "amount": applied["total"]}]}, format="json")

    body = api.get("/api/v1/reports/book/cupons/").json()
    row = next(r for r in next(b for b in body["blocks"] if b["type"] == "compare")["rows"] if r["label"] == "VOLTE10")
    overview = api.get("/api/v1/reports/book/descontos/").json()

    assert row["uses"] == 1 and Decimal(row["given"]) == Decimal("10.00")
    assert Decimal(row["margin_pct"]) == Decimal("55.6")  # (90 − 40) / 90
    assert row["verdict"] in {"Fideliza e dá lucro", "Dá lucro, não fideliza"}
    assert Decimal(overview["kpis"][2]["value"]) == Decimal("10.00")  # em cupons
